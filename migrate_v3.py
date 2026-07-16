#!/usr/bin/env python3
"""
Migration V3 de pp1_collection.db (PP1/Photomaton)
- Ajoute artworks.thumbnail_url (NULL par défaut)
- Normalise artworks.media_url : minuscules, [_ .] -> '-', seuls [a-z0-9-] conservés
  (les extensions sont ajoutées ensuite par sync_images_local.py, exécuté sur la
   machine qui possède le dossier d'images)
- Crée le référentiel `emotions` (emotion, name_fr, name_nl, name_en)
- Ajoute artworks.keywords_fr / keywords_nl / keywords_en
- Ré-importe TOUTES les traductions corrigées depuis review_translations.xlsx :
  descriptions, questions, types d'objets, transcriptions (translation_*),
  explications (explanation_*), témoignages enregistrés (content_*)
Usage: python migrate_v3.py <db_path> <review_translations.xlsx>
"""
import re
import sys
import sqlite3
from openpyxl import load_workbook

DB = sys.argv[1] if len(sys.argv) > 1 else 'pp1_collection.db'
XLSX = sys.argv[2] if len(sys.argv) > 2 else 'review_translations.xlsx'


def normalize_media_name(name: str) -> str:
    """IMadeYou_01_053 -> imadeyou-01-053 ; même algo que sync_images_local.py"""
    s = name.strip().lower()
    s = re.sub(r'[_.\s]+', '-', s)      # _ . espaces -> -
    s = re.sub(r'[^a-z0-9-]', '', s)     # tout le reste supprimé (sauf -)
    s = re.sub(r'-{2,}', '-', s).strip('-')
    return s


def normalize_media_url(value):
    if not value:
        return value
    parts = [p.strip() for p in str(value).split(',')]
    parts = [normalize_media_name(p) for p in parts if p.strip()]
    return ', '.join(parts)


def main():
    con = sqlite3.connect(DB)
    cur = con.cursor()
    wb = load_workbook(XLSX, data_only=True)
    report = []

    def col_exists(table, col):
        return col in [r[1] for r in cur.execute(f'PRAGMA table_info({table})')]

    # ── 1. thumbnail_url ────────────────────────────────────────────────
    if not col_exists('artworks', 'thumbnail_url'):
        cur.execute('ALTER TABLE artworks ADD COLUMN thumbnail_url TEXT')
        report.append('artworks.thumbnail_url ajoutée (NULL)')
    else:
        report.append('artworks.thumbnail_url déjà présente')

    # ── 2. normalisation media_url ──────────────────────────────────────
    changed = 0
    for aid, mu in cur.execute('SELECT id, media_url FROM artworks WHERE media_url IS NOT NULL').fetchall():
        norm = normalize_media_url(mu)
        if norm != mu:
            cur.execute('UPDATE artworks SET media_url=? WHERE id=?', (norm, aid))
            changed += 1
    report.append(f'media_url normalisés : {changed} lignes modifiées')

    # ── 3. référentiel emotions ─────────────────────────────────────────
    cur.execute('''CREATE TABLE IF NOT EXISTS emotions (
        emotion TEXT PRIMARY KEY,
        name_fr TEXT NOT NULL,
        name_nl TEXT NOT NULL,
        name_en TEXT NOT NULL)''')
    ws = wb['Emotions']
    n = 0
    for src, fr, nl, en in ws.iter_rows(min_row=2, values_only=True):
        if not src:
            continue
        key = str(src).strip().lower()
        cur.execute('INSERT OR REPLACE INTO emotions VALUES (?,?,?,?)',
                    (key, str(fr).strip(), str(nl).strip(), str(en).strip()))
        n += 1
    report.append(f'référentiel emotions : {n} entrées (feuille)')
    # Typos dans artwork_emotions -> normalisation des données
    fixed = cur.execute("UPDATE artwork_emotions SET emotion='submission' WHERE lower(trim(emotion))='submision'").rowcount
    fixed += cur.execute("UPDATE artwork_emotions SET emotion='disapproval' WHERE lower(trim(emotion))='dissaproval'").rowcount
    report.append(f'typos corrigées dans artwork_emotions (submision/dissaproval) : {fixed} lignes')
    # Complément Plutchik pour les émotions utilisées mais absentes de la feuille (A VALIDER)
    COMPLEMENT = {
        'aggressiveness': ('Agressivité', 'Agressiviteit', 'Aggressiveness'),
        'amazement':      ('Stupéfaction', 'Verbazing', 'Amazement'),
        'anger':          ('Colère', 'Woede', 'Anger'),
        'annoyance':      ('Agacement', 'Ergernis', 'Annoyance'),
        'contempt':       ('Mépris', 'Minachting', 'Contempt'),
        'disapproval':    ('Désapprobation', 'Afkeuring', 'Disapproval'),
        'interest':       ('Intérêt', 'Interesse', 'Interest'),
        'pensiveness':    ('Rêverie', 'Bedachtzaamheid', 'Pensiveness'),
        'remorse':        ('Remords', 'Wroeging', 'Remorse'),
        'sadness':        ('Tristesse', 'Verdriet', 'Sadness'),
        'surprise':       ('Surprise', 'Verrassing', 'Surprise'),
    }
    added = 0
    for key, (fr, nl, en) in COMPLEMENT.items():
        added += cur.execute('INSERT OR IGNORE INTO emotions VALUES (?,?,?,?)', (key, fr, nl, en)).rowcount
    report.append(f'référentiel emotions : {added} entrées complétées (Plutchik, à valider)')
    # contrôle de couverture : chaque emotion utilisée doit exister dans le référentiel
    missing = cur.execute('''SELECT DISTINCT lower(trim(ae.emotion)) FROM artwork_emotions ae
        LEFT JOIN emotions e ON e.emotion = lower(trim(ae.emotion))
        WHERE e.emotion IS NULL''').fetchall()
    report.append(f'emotions utilisées sans traduction : {[m[0] for m in missing] or "aucune"}')

    # ── 4. keywords i18n ────────────────────────────────────────────────
    for c in ('keywords_fr', 'keywords_nl', 'keywords_en'):
        if not col_exists('artworks', c):
            cur.execute(f'ALTER TABLE artworks ADD COLUMN {c} TEXT')
    ws = wb['Description - Key words']
    n = miss = 0
    for i, _src, fr, nl, en in ws.iter_rows(min_row=2, values_only=True):
        if i is None:
            continue
        r = cur.execute('UPDATE artworks SET keywords_fr=?, keywords_nl=?, keywords_en=? WHERE id=?',
                        (fr, nl, en, int(i)))
        if r.rowcount:
            n += 1
        else:
            miss += 1
    report.append(f'keywords i18n : {n} artworks mis à jour, {miss} ids sans correspondance')

    # ── 5. ré-import des traductions corrigées ──────────────────────────
    # 5a. descriptions
    ws = wb['Descriptions']
    n = 0
    for i, _src, fr, nl, en in ws.iter_rows(min_row=2, values_only=True):
        if i is None:
            continue
        n += cur.execute('UPDATE artworks SET description_fr=?, description_nl=?, description_en=? WHERE id=?',
                         (fr, nl, en, int(i))).rowcount
    report.append(f'descriptions ré-importées : {n}')

    # 5b. questions
    ws = wb['Questions']
    n = 0
    for i, _src, fr, nl, en in ws.iter_rows(min_row=2, values_only=True):
        if i is None:
            continue
        n += cur.execute('UPDATE questions SET content_fr=?, content_nl=?, content_en=? WHERE id=?',
                         (fr, nl, en, int(i))).rowcount
    report.append(f'questions ré-importées : {n}')

    # 5c. types d'objets
    ws = wb["Types d'objets"]
    n = 0
    for i, _src, fr, nl, en in ws.iter_rows(min_row=2, values_only=True):
        if i is None:
            continue
        n += cur.execute('UPDATE types_of_object SET name_fr=?, name_nl=?, name_en=? WHERE id=?',
                         (fr, nl, en, int(i))).rowcount
    report.append(f"types d'objets ré-importés : {n}")

    # 5d. transcriptions (translation_*) — clé = artwork_id
    ws = wb['Transcriptions']
    n = 0
    for i, _src, fr, nl, en in ws.iter_rows(min_row=2, values_only=True):
        if i is None:
            continue
        n += cur.execute('''UPDATE transcriptions SET translation_fr=?, translation_nl=?, translation_en=?
                            WHERE artwork_id=?''', (fr, nl, en, int(i))).rowcount
    report.append(f'transcriptions (translation) ré-importées : {n}')

    # 5e. explications (explanation_*) — fr = source
    ws = wb['Transcriptions Explications']
    n = 0
    for i, _src, fr, nl, en in ws.iter_rows(min_row=2, values_only=True):
        if i is None:
            continue
        n += cur.execute('''UPDATE transcriptions SET explanation_fr=?, explanation_nl=?, explanation_en=?
                            WHERE artwork_id=?''', (fr, nl, en, int(i))).rowcount
    report.append(f'transcriptions (explanation) ré-importées : {n}')

    # 5f. témoignages enregistrés — clé = source_file
    ws = wb['Témoignages']
    n = miss = 0
    for q, _ville, _date, _sl, fr, nl, en, sf in ws.iter_rows(min_row=2, values_only=True):
        if not sf:
            continue
        m = re.search(r'(\d+)', str(q or ''))
        qid = int(m.group(1)) if m else None
        r = cur.execute('''UPDATE recorded_testimonies SET content_fr=?, content_nl=?, content_en=?
                           WHERE source_file=? AND question_id=?''',
                        (fr, nl, en, str(sf).strip(), qid))
        if r.rowcount == 1:
            n += 1
        else:
            miss += 1
    report.append(f'témoignages ré-importés : {n}, sans correspondance exacte : {miss}')

    con.commit()
    con.close()
    print('\n'.join('  • ' + r for r in report))


if __name__ == '__main__':
    main()

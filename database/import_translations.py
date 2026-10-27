#!/usr/bin/env python3
"""
Loads the texts of the review workbook (review_translations.xlsx) into the database.

The review workbook is the single source of truth for every text shown to
visitors in FR / NL / EN. Run this script as often as needed:
  - after translate_collection.py, to load the machine translations;
  - after each review round, to load the reviewers' corrections.

Sheets read (a missing sheet is simply skipped):
  Types d'objets, Questions, Descriptions, Description - Key words, Emotions,
  Transcriptions, Transcriptions Explications, Témoignages, Titres.

An empty cell never overwrites a text already in the database: clearing a
cell in the workbook deletes nothing.

Usage:
    python import_translations.py <db> <review_translations.xlsx>
Idempotent: running it again on the same workbook changes nothing.

(Formerly migrate_v3.py, without its one-off operations, which were applied
long ago.)
"""
import sqlite3
import sys

from openpyxl import load_workbook


def clean(value):
    """The cell text as is, or None if the cell is empty."""
    if value is None or not str(value).strip():
        return None
    return value if isinstance(value, str) else str(value)


def as_id(value):
    """Row id as an int, or None for a row without a numeric id."""
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def import_lang_sheet(cur, wb, sheet, sql, report, label, source_sql=None):
    """Sheets laid out as id | source | fr | nl | en.

    source_sql: query returning the current source text in the database. If
    the workbook's "source" cell no longer matches it (the museum corrected
    its text after the translation), the row is skipped: translate_collection.py
    must run first, otherwise old translations would overwrite the new text.
    """
    if sheet not in wb.sheetnames:
        report.append(f'{label}: sheet "{sheet}" not found, skipped')
        return
    n = miss = 0
    stale = []
    for row in wb[sheet].iter_rows(min_row=2, values_only=True):
        i, src, fr, nl, en = (list(row or []) + [None] * 5)[:5]
        i = as_id(i)
        if i is None:
            continue
        if source_sql:
            current = cur.execute(source_sql, (i,)).fetchone()
            if current and str(current[0] or '').strip() != str(src or '').strip():
                stale.append(i)
                continue
        r = cur.execute(sql, (clean(fr), clean(nl), clean(en), i))
        if r.rowcount:
            n += 1
        else:
            miss += 1
    report.append(f'{label}: {n} row(s) loaded'
                  + (f', {miss} id(s) not in the database' if miss else ''))
    if stale:
        report.append(f'⚠ {label}: {len(stale)} row(s) skipped because the source text changed '
                      f'after translation {stale[:10]} → run translate_collection.py first')


def main(db_path, xlsx_path):
    con = sqlite3.connect(db_path)
    cur = con.cursor()
    wb = load_workbook(xlsx_path, data_only=True)
    report = []

    # The workbook's "source" column (grey cells) is a reference for reviewers:
    # it is never written back to the database.
    import_lang_sheet(cur, wb, "Types d'objets", '''
        UPDATE types_of_object SET name_fr = COALESCE(?, name_fr), name_nl = COALESCE(?, name_nl),
            name_en = COALESCE(?, name_en)
        WHERE id = ?''', report, 'object types', 'SELECT name FROM types_of_object WHERE id = ?')

    import_lang_sheet(cur, wb, 'Questions', '''
        UPDATE questions SET content_fr = COALESCE(?, content_fr), content_nl = COALESCE(?, content_nl),
            content_en = COALESCE(?, content_en)
        WHERE id = ?''', report, 'questions')

    import_lang_sheet(cur, wb, 'Descriptions', '''
        UPDATE artworks SET description_fr = COALESCE(?, description_fr),
            description_nl = COALESCE(?, description_nl), description_en = COALESCE(?, description_en)
        WHERE id = ?''', report, 'descriptions', 'SELECT description FROM artworks WHERE id = ?')

    import_lang_sheet(cur, wb, 'Description - Key words', '''
        UPDATE artworks SET keywords_fr = COALESCE(?, keywords_fr), keywords_nl = COALESCE(?, keywords_nl),
            keywords_en = COALESCE(?, keywords_en)
        WHERE id = ?''', report, 'keywords', 'SELECT keywords FROM artworks WHERE id = ?')

    import_lang_sheet(cur, wb, 'Titres', '''
        UPDATE artworks SET title_fr = COALESCE(?, title_fr), title_nl = COALESCE(?, title_nl),
            title_en = COALESCE(?, title_en)
        WHERE id = ?''', report, 'titles', 'SELECT title FROM artworks WHERE id = ?')

    # Transcriptions: key = artwork_id
    import_lang_sheet(cur, wb, 'Transcriptions', '''
        UPDATE transcriptions SET translation_fr = COALESCE(?, translation_fr),
            translation_nl = COALESCE(?, translation_nl), translation_en = COALESCE(?, translation_en)
        WHERE artwork_id = ?''', report, 'transcriptions',
        'SELECT coalesce(translation, transcription) FROM transcriptions WHERE artwork_id = ?')

    import_lang_sheet(cur, wb, 'Transcriptions Explications', '''
        UPDATE transcriptions SET explanation_fr = COALESCE(?, explanation_fr),
            explanation_nl = COALESCE(?, explanation_nl), explanation_en = COALESCE(?, explanation_en)
        WHERE artwork_id = ?''', report, 'transcription explanations',
        'SELECT explanation FROM transcriptions WHERE artwork_id = ?')

    # Emotions: key -> labels (layout source | fr | nl | en)
    if 'Emotions' in wb.sheetnames:
        n = 0
        for row in wb['Emotions'].iter_rows(min_row=2, values_only=True):
            src, fr, nl, en = (list(row) + [None] * 4)[:4]
            key = clean(src) and clean(src).strip()
            if not key or not (clean(fr) and clean(nl) and clean(en)):
                continue
            cur.execute('INSERT OR REPLACE INTO emotions (emotion, name_fr, name_nl, name_en) VALUES (?,?,?,?)',
                        (key.lower(), clean(fr).strip(), clean(nl).strip(), clean(en).strip()))
            n += 1
        report.append(f'emotions: {n} row(s) loaded')
    missing = [m[0] for m in cur.execute('''
        SELECT DISTINCT lower(trim(ae.emotion)) FROM artwork_emotions ae
        LEFT JOIN emotions e ON e.emotion = lower(trim(ae.emotion)) WHERE e.emotion IS NULL''')]
    if missing:
        report.append(f'⚠ emotions used by artworks but not translated: {missing}')

    # Recorded interviews: column H holds the testimony id in the database
    # (interviewees are no longer identified by their first name).
    if 'Témoignages' in wb.sheetnames:
        ws = wb['Témoignages']
        if str(ws.cell(1, 8).value or '').strip().lower() != 'id':
            raise SystemExit('Sheet "Témoignages": column H must be named "id" (testimony id). '
                             'Old layout with "fichier_source": use the up-to-date workbook from the repository.')
        n = miss = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            _q, _city, date, src_lang, fr, nl, en, tid = (list(row) + [None] * 8)[:8]
            if tid is None:
                continue
            lang = (clean(src_lang) or 'fr').lower()
            content = clean({'fr': fr, 'nl': nl, 'en': en}.get(lang, fr))
            created = date.date().isoformat() if hasattr(date, 'date') else (str(date)[:10] if date else None)
            r = cur.execute('''
                UPDATE testimonies SET content = COALESCE(?, content),
                    content_fr = COALESCE(?, content_fr), content_nl = COALESCE(?, content_nl),
                    content_en = COALESCE(?, content_en), created_at = COALESCE(?, created_at)
                WHERE id = ? AND visitor_id IS NULL''',
                (content, clean(fr), clean(nl), clean(en), created, int(tid)))
            if r.rowcount:
                n += 1
            else:
                miss += 1
        report.append(f'interviews: {n} row(s) loaded' + (f', {miss} id(s) not in the database' if miss else ''))

    # Final check: artworks still incomplete
    incomplete = [r[0] for r in cur.execute('''
        SELECT id FROM artworks
        WHERE description_fr IS NULL OR description_nl IS NULL OR description_en IS NULL
           OR title_fr IS NULL OR title_nl IS NULL OR title_en IS NULL''')]
    con.commit()
    con.close()
    print('\n'.join('  • ' + r for r in report))
    if incomplete:
        print(f'  ⚠ {len(incomplete)} artwork(s) still missing a title or description in one of the 3 languages: '
              f'{incomplete[:15]}{" …" if len(incomplete) > 15 else ""} → run translate_collection.py')
    else:
        print(f'  → next step: python compute_embeddings.py {db_path}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    from pathlib import Path
    for path, what in ((sys.argv[1], 'database'), (sys.argv[2], 'workbook')):
        if not Path(path).is_file():
            sys.exit(f'{what} not found: {path} (run the command from the database/ folder)')
    main(sys.argv[1], sys.argv[2])

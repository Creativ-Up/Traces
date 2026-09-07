#!/usr/bin/env python3
"""
Migration V10 de pp1_collection.db (PP1/Photomaton)
Intègre les collections des deux nouveaux partenaires INTERREG :
  - Huis van Alijn : 99 œuvres (ids 121–219), textes en néerlandais
  - MUMONS         : 15 œuvres (ids 220–234), textes en français

Source : PP1-Collection_Database.xlsx consolidé (feuille Database, lignes 121+).
Les médias doivent avoir été préparés au préalable avec prepare_media.py
(TIF/PNG -> JPG, MOV/NEF écartés, noms normalisés) et copiés dans assets/.

Remplit : museum_id, origin, type_of_object_id, date_period + années,
title (+ title_nl ou title_fr selon la langue source), author_name,
storage_place, keywords (+ keywords_nl/fr), description (+ description_nl/fr),
question_id, media_url, thumbnail_url, emotions, et la table artwork_emotions.
Les traductions manquantes restent NULL : elles arrivent par la feuille
« Titres » / « Descriptions » du classeur de relecture via migrate_v3.py.

Usage : python migrate_v10.py <db_path> <PP1-Collection_Database.xlsx>
Idempotent : les œuvres déjà présentes sont mises à jour, pas dupliquées.
"""
import re
import sys
import sqlite3

from openpyxl import load_workbook

DB = sys.argv[1] if len(sys.argv) > 1 else 'pp1_collection.db'
XLSX = sys.argv[2] if len(sys.argv) > 2 else 'PP1-Collection_Database.xlsx'

FIRST_NEW_ID = 121                      # les 116 œuvres historiques vont jusqu'à 120
LANG_BY_ORIGIN = {'huis van alijn': 'nl', 'mumons': 'fr'}

IMG_EXTS = ('.jpg', '.jpeg', '.png', '.webp', '.gif', '.tif', '.tiff', '.bmp', '.avif')
VIDEO_EXTS = ('.mp4', '.webm')
# extensions réelles des médias préparés : tout devient .jpg sauf les vidéos
VIDEO_STEMS_HINT = ('vi-', 'f00529')    # préfixes des fichiers vidéo livrés


def normalize_media_name(name: str) -> str:
    """Identique à migrate_v3.normalize_media_name (extension préservée)."""
    s = name.strip().lower()
    ext = ''
    for e in IMG_EXTS + VIDEO_EXTS:
        if s.endswith(e):
            s, ext = s[:-len(e)], e
            break
    s = re.sub(r'[_.\s]+', '-', s)
    s = re.sub(r'[^a-z0-9-]', '', s)
    s = re.sub(r'-{2,}', '-', s).strip('-')
    return s + ext


def media_filename(ref: str) -> str:
    """Réf. PHOTO de l'Excel -> nom du fichier préparé (.mp4 pour les vidéos, .jpg sinon)."""
    stem = normalize_media_name(ref)
    for e in IMG_EXTS + VIDEO_EXTS:
        if stem.endswith(e):
            stem = stem[:-len(e)]
            break
    ext = '.mp4' if stem.startswith(VIDEO_STEMS_HINT) else '.jpg'
    return stem + ext


def parse_years(value):
    """'1870-1880' -> (1870, 1880) ; 'ca. 1935' -> (1935, 1935) ;
    'first half of the 20th century' -> (1900, 1950)."""
    if value is None:
        return None, None
    s = str(value).lower()
    years = re.findall(r'(1[6-9]\d{2}|20\d{2})', s)
    if years:
        return int(years[0]), int(years[-1])
    m = re.search(r'(\d{1,2})(?:st|nd|rd|th)\s+century', s)
    if m:
        start = (int(m.group(1)) - 1) * 100
        if 'first half' in s:
            return start, start + 50
        if 'second half' in s:
            return start + 50, start + 99
        if 'early' in s:
            return start, start + 30
        if 'late' in s:
            return start + 70, start + 99
        return start, start + 99
    return None, None


def main():
    con = sqlite3.connect(DB)
    cur = con.cursor()
    report = []

    types = {str(n).strip().lower(): i for i, n in cur.execute('SELECT id, name FROM types_of_object')}
    questions = {}
    for qid, txt in cur.execute('SELECT id, COALESCE(content_en, content) FROM questions'):
        questions[re.sub(r'[^a-z0-9 ]', '', str(txt).lower()).strip()] = qid
    known_emotions = {e.lower() for (e,) in cur.execute('SELECT emotion FROM emotions')}

    ws = load_workbook(XLSX, data_only=True)['Database']
    inserted = updated = 0
    warn = []

    for r in range(3, ws.max_row + 1):
        raw_id = ws.cell(r, 1).value
        if raw_id is None or not str(raw_id).strip():
            continue
        aid = int(raw_id)
        if aid < FIRST_NEW_ID:
            continue                                  # œuvres historiques : intactes

        origin = str(ws.cell(r, 3).value or '').strip()
        lang = LANG_BY_ORIGIN.get(origin.lower(), 'fr')

        type_name = str(ws.cell(r, 4).value or '').strip()
        type_id = types.get(type_name.lower())
        if type_id is None:
            warn.append(f'œuvre {aid} : type inconnu « {type_name} »')

        q_raw = str(ws.cell(r, 16).value or '')
        q_id = questions.get(re.sub(r'[^a-z0-9 ]', '', q_raw.lower()).strip())
        if q_id is None and q_raw.strip():
            warn.append(f'œuvre {aid} : question non reconnue')

        date_period = ws.cell(r, 5).value
        y_min, y_max = parse_years(date_period)
        # Excel renvoie certaines dates HVA comme datetime ('1962-12-24 00:00:00') :
        # seule l'annee est conservee pour l'affichage (choix valide : le mois et le
        # jour ne sont ni fiables ni utiles), les annees servent au matching.
        if hasattr(date_period, 'year') or re.match(r'^\d{4}-\d{2}-\d{2}[ T]', str(date_period or '')):
            date_period = str(y_min) if y_min else None

        photos = [media_filename(p) for p in re.split(r'[;,:\n]+', str(ws.cell(r, 17).value or '')) if p.strip()]
        media_url = ', '.join(photos) if photos else None
        thumb = next((p for p in photos if p.endswith('.jpg')), None)  # jamais une vidéo en vignette
        if thumb is None and photos:
            # œuvre uniquement vidéo : vignette extraite par prepare_media.py
            thumb = photos[0].rsplit('.', 1)[0] + '-thumb.jpg'

        emotions_raw = [e.strip() for e in re.split(r'[;,\n]+', str(ws.cell(r, 14).value or '')) if e.strip()]
        emotions_txt = ', '.join(emotions_raw) if emotions_raw else None

        title = ws.cell(r, 6).value
        keywords = ws.cell(r, 12).value
        description = ws.cell(r, 15).value

        row = {
            'id': aid, 'museum_id': str(ws.cell(r, 2).value or '').strip() or None,
            'origin': origin or None, 'type_of_object_id': type_id,
            'date_period': str(date_period).strip() if date_period is not None else None,
            'date_year_min': y_min, 'date_year_max': y_max,
            'title': title, f'title_{lang}': title,
            'author_name': ws.cell(r, 8).value, 'storage_place': ws.cell(r, 10).value,
            'keywords': keywords, f'keywords_{lang}': keywords,
            'description': description, f'description_{lang}': description,
            'question_id': q_id, 'media_url': media_url, 'thumbnail_url': thumb,
            'emotions': emotions_txt, 'popularity': 0,
        }

        exists = cur.execute('SELECT 1 FROM artworks WHERE id=?', (aid,)).fetchone()
        cols = list(row)
        if exists:
            cur.execute(f"UPDATE artworks SET {', '.join(c + '=?' for c in cols if c != 'id')} WHERE id=?",
                        [row[c] for c in cols if c != 'id'] + [aid])
            updated += 1
        else:
            cur.execute(f"INSERT INTO artworks ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                        [row[c] for c in cols])
            inserted += 1

        cur.execute('DELETE FROM artwork_emotions WHERE artwork_id=?', (aid,))
        for e in emotions_raw:
            key = e.lower()
            if key not in known_emotions:
                warn.append(f'œuvre {aid} : émotion hors référentiel « {e} »')
            cur.execute('INSERT INTO artwork_emotions (artwork_id, emotion) VALUES (?, ?)', (aid, key))

    report.append(f'œuvres insérées : {inserted}, mises à jour : {updated}')
    total = cur.execute('SELECT COUNT(*) FROM artworks').fetchone()[0]
    report.append(f'collection : {total} œuvres')
    for o, n in cur.execute('SELECT origin, COUNT(*) FROM artworks GROUP BY origin ORDER BY 2 DESC'):
        report.append(f'  {o} : {n}')
    missing_tr = cur.execute("""SELECT COUNT(*) FROM artworks
                                WHERE id >= ? AND (description_fr IS NULL OR description_nl IS NULL
                                                   OR description_en IS NULL)""", (FIRST_NEW_ID,)).fetchone()[0]
    report.append(f'œuvres en attente de traduction : {missing_tr} (feuilles du classeur de relecture)')
    for w in dict.fromkeys(warn):
        report.append(f'⚠ {w}')

    con.commit()
    con.close()
    print('\n'.join('  • ' + x for x in report))


if __name__ == '__main__':
    main()

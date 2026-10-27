#!/usr/bin/env python3
"""
Imports a museum's collection into pp1_collection.db, from a spreadsheet in the
TRACES format (template_collection.xlsx). The historical workbook
PP1-Collection_Database.xlsx has the same format but must only be read: the
database has changed since (--update is refused with that file).

For each artwork in the spreadsheet:
  - the artwork is created (or updated with --update) in `artworks`;
  - its object type, question and emotions are linked to the database
    reference lists (an unknown object type is created, to be translated);
  - its media files are looked up in assets/ (.jpg images, .mp4 videos);
  - its transcriptions, if any ("Transcriptions" sheet), are imported.

Texts are stored in their original language. The translations into the two
other languages are prepared next by translate_collection.py.

How artworks are identified:
  - "Database ID" filled in: it is the TRACES id of the artwork;
  - otherwise the artwork is recognised by the pair (Origin, Museum ID) and
    receives a new id when created. Running the import again therefore never
    creates duplicates.
An artwork already in the database is never changed without --update.

Usage:
    python import_collection.py <db> <spreadsheet.xlsx> --assets ../assets
           [--origin "Museum name"] [--lang fr|nl|en] [--update] [--dry-run]

  --assets   folder of the prepared media (prepare_media.py), to check that
             every file named in the spreadsheet exists
  --origin   only import the rows of this museum
  --lang     language of the texts when the "Language" column is empty
  --update   update artworks already in the database (an unrecognised value:
             question, type, emotion, missing media, never replaces the current one)
  --dry-run  check everything and print the report, without writing anything
"""
from __future__ import annotations

import argparse
import re
import sqlite3
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

LANGS = ('fr', 'nl', 'en')
LANG_ALIASES = {
    'fr': 'fr', 'francais': 'fr', 'french': 'fr', 'frans': 'fr',
    'nl': 'nl', 'nederlands': 'nl', 'neerlandais': 'nl', 'dutch': 'nl', 'vlaams': 'nl',
    'en': 'en', 'anglais': 'en', 'english': 'en', 'engels': 'en',
}
IMAGE_EXTS = ('.jpg', '.jpeg', '.png', '.webp', '.gif', '.tif', '.tiff', '.bmp', '.avif')
VIDEO_EXTS = ('.mp4', '.webm', '.mov')
KNOWN_EXTS = IMAGE_EXTS + VIDEO_EXTS

# Accepted column headers (compared without accents, case or punctuation)
FIELDS = {
    'id': ('databaseid', 'idtraces'),
    'museum_id': ('museumid', 'inventorynumber', 'numerodinventaire', 'inventaire', 'inventarisnummer'),
    'origin': ('origin', 'institution', 'musee', 'museum'),
    'type': ('typeofobject', 'typedobjet', 'type'),
    'date': ('dateperiod', 'date', 'period', 'periode', 'datering'),
    'title': ('title', 'titre', 'titel'),
    'author': ('author', 'auteur', 'artist', 'artiste', 'maker'),
    'author_name': ('name', 'nom', 'names', 'personnes'),
    'storage': ('storageplace', 'lieudeconservation', 'bewaarplaats'),
    'keywords': ('descriptionkeywords', 'keywords', 'motscles', 'trefwoorden'),
    'emotions': ('pulchikswheelmatch', 'plutchikswheelmatch', 'plutchik', 'emotions', 'emoties'),
    'description': ('explanation', 'description', 'beschrijving'),
    'question': ('question', 'vraag'),
    'media': ('photo', 'photos', 'media', 'medias', 'files', 'fichiers', 'bestanden'),
    'lang': ('language', 'langue', 'taal', 'languedestextes'),
}
TRANSCRIPTION_FIELDS = {
    'id': FIELDS['id'],
    'museum_id': FIELDS['museum_id'],
    'transcription': ('transcriptions', 'transcription', 'transcriptie'),
    'explanation': ('explanation', 'explanations', 'explication', 'explications', 'uitleg'),
    'translation': ('translation', 'traduction', 'vertaling'),
}


def norm(text) -> str:
    s = unicodedata.normalize('NFKD', str(text or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]', '', s.lower())


def text(value):
    """Cell value as text, None if empty."""
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def fold(s: str) -> str:
    """Removes accents (é -> e), so that a file name typed in Excel and the same
    name written by macOS (decomposed accents) give the same result."""
    return unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode()


def normalize_media_name(name: str) -> str:
    """TRACES media naming convention (same as prepare_media.py):
    IMadeYou_01_053.JPG -> imadeyou-01-053.jpg (extension kept)."""
    s = fold(name.strip().lower())
    ext = ''
    for e in KNOWN_EXTS:
        if s.endswith(e):
            s, ext = s[:-len(e)], e
            break
    s = re.sub(r'[_.\s]+', '-', s)
    s = re.sub(r'[^a-z0-9-]', '', s)
    s = re.sub(r'-{2,}', '-', s).strip('-')
    return s + ext


ROMAN = {'i': 1, 'v': 5, 'x': 10, 'l': 50, 'c': 100}


def roman_to_int(s):
    total = 0
    for a, b in zip(s, s[1:] + ' '):
        v = ROMAN[a]
        total += -v if b in ROMAN and ROMAN[b] > v else v
    return total


def parse_years(value):
    """(first year, last year) of a date or period, (None, None) if unreadable.
    '1870-1880' -> (1870, 1880) ; '1914-18' -> (1914, 1918) ; 'ca. 1935' -> (1935, 1935) ;
    '19e siècle' or 'XIXe siècle' -> (1800, 1899) ; 'first half of the 20th century' -> (1900, 1950)."""
    if value is None:
        return None, None
    if hasattr(value, 'year'):
        return value.year, value.year
    s = fold(str(value).lower())
    m = re.search(r'(?<!\d)(1\d{3}|20\d{2})\s*[-/]\s*(\d{2})(?!\d)', s)       # 1914-18, 1950-60
    if m:
        a, b = int(m.group(1)), int(m.group(1)[:2] + m.group(2))
        return min(a, b), max(a, b)
    years = [int(y) for y in re.findall(r'(?<!\d)(1\d{3}|20\d{2})(?!\d)', s)]
    if years:
        return min(years), max(years)                                          # 1890-1880 -> (1880, 1890)
    m = re.search(r'(?<![a-z0-9])(\d{1,2}|[ivxlc]+)\s*(?:st|nd|rd|th|e|eme|er|de|ste)?\s+'
                  r'(?:century|siecle|eeuw)', s)
    if m:
        n = int(m.group(1)) if m.group(1).isdigit() else roman_to_int(m.group(1))
        start = (n - 1) * 100
        if any(w in s for w in ('first half', 'premiere moitie', 'eerste helft')):
            return start, start + 50
        if any(w in s for w in ('second half', 'seconde moitie', 'deuxieme moitie', 'tweede helft')):
            return start + 50, start + 99
        if any(w in s for w in ('early', 'debut', 'begin')):
            return start, start + 30
        if any(w in s for w in ('late', 'fin ', 'eind')):
            return start + 70, start + 99
        return start, start + 99
    return None, None


def find_header(ws, fields, required):
    """Finds the header row (among the first 6) and the column of each field."""
    for r in range(1, min(ws.max_row, 6) + 1):
        cols = {}
        for c in range(1, ws.max_column + 1):
            key = norm(ws.cell(r, c).value)
            for field, aliases in fields.items():
                if key in aliases and field not in cols:
                    cols[field] = c
        if all(f in cols for f in required):
            return r, cols
    return None, None


def pick_sheet(wb):
    for name in ('Artworks', 'Œuvres', 'Oeuvres', 'Database', 'Collection'):
        if name in wb.sheetnames:
            return wb[name]
    for ws in wb.worksheets:
        if find_header(ws, FIELDS, ('origin', 'description'))[0]:
            return ws
    raise SystemExit('no artworks sheet found (expected headers "Origin" and "Explanation")')


class Importer:
    def __init__(self, con, args):
        self.con, self.cur, self.args = con, con.cursor(), args
        self.report, self.warn = [], []
        self.assets = None
        if args.assets:
            folder = Path(args.assets)
            if not folder.is_dir():
                raise SystemExit(f'media folder not found: {folder}')
            self.assets = {p.name for p in folder.iterdir() if p.is_file()}
        self.next_id = (self.cur.execute('SELECT COALESCE(MAX(id), 0) FROM artworks').fetchone()[0]) + 1
        self.types = {}
        for tid, *names in self.cur.execute('SELECT id, name, name_fr, name_nl, name_en FROM types_of_object'):
            for n in names:
                if n:
                    self.types.setdefault(norm(n), tid)
        self.questions = {}
        for qid, *contents in self.cur.execute('SELECT id, content, content_fr, content_nl, content_en FROM questions'):
            for c in contents:
                if c:
                    self.questions[norm(c)] = qid
        self.question_ids = {q for q in self.questions.values()}
        self.emotions = {}
        for key, *names in self.cur.execute('SELECT emotion, name_fr, name_nl, name_en FROM emotions'):
            self.emotions[norm(key)] = key
            for n in names:
                if n:
                    self.emotions.setdefault(norm(n), key)
        self.new_types = []
        self.missing = defaultdict(list)   # empty column -> artworks concerned
        self.mapping = []                  # (Museum ID, new id) of created artworks
        self.touched = {}          # museum_id / id -> artwork id, for the transcriptions

    # ── matching against the reference lists ─────────────────────────────
    # Each resolver returns (value, ok): ok=False means a value was given but
    # not recognised. On update, the current value is then kept rather than
    # overwritten.
    def resolve_lang(self, value, aid):
        if value:
            lang = LANG_ALIASES.get(norm(value))
            if lang:
                return lang
            self.warn.append(f'artwork {aid}: unknown language "{value}" (fr, nl or en)')
        return self.args.lang

    def resolve_type(self, name, lang, aid, create):
        if not name:
            return None, False
        tid = self.types.get(norm(name))
        if tid is not None:
            return tid, True
        if not create:
            self.warn.append(f'artwork {aid}: unknown object type "{name}", current type kept')
            return None, False
        self.cur.execute('INSERT INTO types_of_object (name) VALUES (?)', (name,))
        tid = self.cur.lastrowid
        if lang:
            self.cur.execute(f'UPDATE types_of_object SET name_{lang} = ? WHERE id = ?', (name, tid))
        self.types[norm(name)] = tid
        self.new_types.append(f'{name} (id {tid})')
        return tid, True

    def resolve_question(self, value, aid):
        raw = text(value)
        if not raw:
            self.missing['Question'].append(aid)
            return None, True
        m = re.fullmatch(r'(?:question|vraag|q)?\s*(?:n\s*°?|°)?\s*(\d+)', raw.strip().lower())
        if m and int(m.group(1)) in self.question_ids:
            return int(m.group(1)), True
        qid = self.questions.get(norm(raw))
        if qid is None:
            self.warn.append(f'artwork {aid}: question not recognised "{raw[:60]}" (expected a number from 1 to 12)')
            return None, False
        return qid, True

    def resolve_emotions(self, value, aid):
        keys, ok = [], True
        for token in re.split(r'[;,\n/]+', str(value or '')):
            if not token.strip():
                continue
            key = self.emotions.get(norm(token))
            if key is None:
                self.warn.append(f'artwork {aid}: emotion "{token.strip()}" is not in the reference list, '
                                 'ignored (see the "Emotions" sheet of the template)')
                ok = False
            elif key not in keys:
                keys.append(key)
        return keys, ok

    def resolve_media(self, value, aid):
        files, ok = [], True
        for token in re.split(r'[;,\n]+', str(value or '')):
            if not token.strip():
                continue
            name = normalize_media_name(token)
            stem, ext = (name.rsplit('.', 1)[0], '.' + name.rsplit('.', 1)[1]) if '.' in name else (name, '')
            found = None
            if self.assets is not None:
                if name in self.assets and ext not in ('.tif', '.tiff', '.png', '.mov', '.jpeg'):
                    found = name
                else:
                    for e in ('.jpg', '.mp4', '.webm', '.webp'):
                        if stem + e in self.assets:
                            found = stem + e
                            break
            if found is None:
                # expected name after prepare_media.py: videos as .mp4, everything else as .jpg
                found = stem + ('.mp4' if ext in VIDEO_EXTS else '.jpg')
                if self.assets is not None:
                    self.warn.append(f'artwork {aid}: media "{token.strip()}" not found in assets/ (expected: {found})')
                    ok = False
            files.append(found)
        images = [f for f in files if not f.endswith(VIDEO_EXTS)]
        thumb = images[0] if images else (files[0].rsplit('.', 1)[0] + '-thumb.jpg' if files else None)
        if thumb and self.assets is not None and thumb not in self.assets:
            self.warn.append(f'artwork {aid}: thumbnail "{thumb}" not found in assets/')
            ok = False
        if not files:
            self.warn.append(f'artwork {aid}: no media')
            ok = False
        return (', '.join(files) or None), thumb, ok

    # ── importing artworks ────────────────────────────────────────────────
    def import_artworks(self, ws):
        hrow, cols = find_header(ws, FIELDS, ('origin', 'description'))
        if not hrow:
            raise SystemExit(f'sheet "{ws.title}": headers "Origin" and "Explanation" not found')
        get = lambda r, f: ws.cell(r, cols[f]).value if f in cols else None
        inserted, updated, skipped = [], 0, 0
        origin_filter = norm(self.args.origin) if self.args.origin else None
        seen = set()

        for r in range(hrow + 1, ws.max_row + 1):
            origin = text(get(r, 'origin'))
            description = text(get(r, 'description'))
            if not origin and not description:
                continue                                        # empty row
            if origin_filter and norm(origin) != origin_filter:
                continue
            raw_id, museum_id = text(get(r, 'id')), text(get(r, 'museum_id'))
            label = f'row {r}'
            if not origin or not description:
                self.warn.append(f'{label}: "Origin" and "Explanation" are required, row skipped')
                continue

            # identification
            if raw_id:
                try:
                    aid = int(float(raw_id))
                except ValueError:
                    self.warn.append(f'{label}: invalid Database ID "{raw_id}", row skipped')
                    continue
                row = self.cur.execute('SELECT origin FROM artworks WHERE id = ?', (aid,)).fetchone()
                if row and norm(row[0]) != norm(origin):
                    self.warn.append(f'{label}: id {aid} already belongs to "{row[0]}", row skipped')
                    continue
                exists = row is not None
                key = ('id', aid)
            elif museum_id:
                # same museum = same name, ignoring case, accents and punctuation
                match = [i for i, o in self.cur.execute('SELECT id, origin FROM artworks WHERE museum_id = ?',
                                                        (museum_id,)) if norm(o) == norm(origin)]
                exists = bool(match)
                aid = match[0] if exists else self.next_id
                key = ('museum_id', norm(origin), museum_id)
            else:
                self.warn.append(f'{label}: "Database ID" or "Museum ID" is required, row skipped')
                continue
            if key in seen:
                self.warn.append(f'{label}: artwork listed twice in the spreadsheet ({museum_id or aid}), row skipped')
                continue
            seen.add(key)
            label = f'artwork {aid}'
            if exists and not self.args.update:
                skipped += 1
                continue
            old = {}
            if exists:
                old = dict(zip(('question_id', 'media_url', 'thumbnail_url', 'type_of_object_id', 'emotions'),
                               self.cur.execute('''SELECT question_id, media_url, thumbnail_url, type_of_object_id,
                                                          emotions FROM artworks WHERE id = ?''', (aid,)).fetchone()))

            lang = self.resolve_lang(text(get(r, 'lang')), aid)
            if not lang:
                self.missing['Language (DeepL will detect it)'].append(aid)
            type_id, type_ok = self.resolve_type(text(get(r, 'type')), lang, aid, create=not exists)
            if not type_ok:
                if not exists:
                    self.warn.append(f'{label}: object type missing, row skipped')
                    continue
                type_id = old['type_of_object_id']
            title = text(get(r, 'title'))
            if not title:
                self.missing['Title'].append(aid)
            keywords = text(get(r, 'keywords'))
            if not keywords:
                self.missing['Keywords'].append(aid)
            date_period = get(r, 'date')
            y_min, y_max = parse_years(date_period)
            if hasattr(date_period, 'year'):          # typed as an Excel date: the year is enough
                date_period = str(date_period.year)
            if date_period is None:
                self.missing['Date /period'].append(aid)
            elif y_min is None:
                self.warn.append(f'{label}: unreadable date "{date_period}" (used to pick related artworks)')
            question_id, q_ok = self.resolve_question(get(r, 'question'), aid)
            media_url, thumb, media_ok = self.resolve_media(get(r, 'media'), aid)
            emotions, emo_ok = self.resolve_emotions(get(r, 'emotions'), aid)
            if not emotions:
                self.warn.append(f'{label}: no recognised emotion (emotions are used to pick related artworks)')

            values = {
                'museum_id': museum_id, 'origin': origin, 'type_of_object_id': type_id,
                'date_period': text(date_period), 'date_year_min': y_min, 'date_year_max': y_max,
                'title': title, 'author': text(get(r, 'author')), 'author_name': text(get(r, 'author_name')),
                'storage_place': text(get(r, 'storage')) or '',        # never NULL: the backend requires it
                'keywords': keywords or '', 'description': description,
                'question_id': question_id, 'media_url': media_url, 'thumbnail_url': thumb,
                'emotions': ', '.join(emotions),                         # never NULL: the backend requires it
            }
            if exists:                 # an unrecognised value never overwrites the current one
                if not q_ok:
                    values['question_id'] = old['question_id']
                if not media_ok:
                    values['media_url'], values['thumbnail_url'] = old['media_url'], old['thumbnail_url']
                    self.warn.append(f'{label}: current media kept')
                if not emo_ok:
                    values['emotions'] = old['emotions']
                    emotions = [e.strip().lower() for e in (old['emotions'] or '').split(',') if e.strip()]
            if lang:                                   # the original text, in its own language
                values.update({f'title_{lang}': title, f'description_{lang}': description,
                               f'keywords_{lang}': keywords})
            if exists:
                sets = ', '.join(f'{c} = ?' for c in values)
                self.cur.execute(f'UPDATE artworks SET {sets} WHERE id = ?', list(values.values()) + [aid])
                updated += 1
            else:
                values.update({'id': aid, 'popularity': 0})
                cols_sql = ', '.join(values)
                self.cur.execute(f'INSERT INTO artworks ({cols_sql}) VALUES ({", ".join("?" * len(values))})',
                                 list(values.values()))
                self.next_id = max(self.next_id, aid + 1)
                inserted.append((museum_id, aid))
            self.cur.execute('DELETE FROM artwork_emotions WHERE artwork_id = ?', (aid,))
            self.cur.executemany('INSERT INTO artwork_emotions (artwork_id, emotion) VALUES (?, ?)',
                                 [(aid, e) for e in emotions])
            self.touched[('id', str(aid))] = (aid, lang)
            if museum_id:
                self.touched[('museum_id', museum_id)] = (aid, lang)

        self.report.append(f'artworks created: {len(inserted)}, updated: {updated}, '
                           f'already in the database and left unchanged: {skipped}'
                           + (' (add --update to change them)' if skipped and not self.args.update else ''))
        if inserted:
            ids = [a for _m, a in inserted]
            self.report.append(f'new TRACES ids: {min(ids)}-{max(ids)} (give them to your reviewers; '
                               f'use --ids {min(ids)}-{max(ids)} with translate_collection.py)')
            self.mapping = inserted
        if self.new_types:
            self.report.append(f'new object types created, to be translated (review the rows with these ids in '
                               f'the "Types d\'objets" sheet): {self.new_types}')
        for col, ids in self.missing.items():
            self.warn.append(f'column "{col}" empty for {len(ids)} artwork(s): {ids[:15]}{" …" if len(ids) > 15 else ""}')

    # ── transcriptions (optional sheet) ─────────────────────────────
    def import_transcriptions(self, wb):
        if 'Transcriptions' not in wb.sheetnames:
            return
        ws = wb['Transcriptions']
        hrow, cols = find_header(ws, TRANSCRIPTION_FIELDS, ('transcription',))
        if not hrow:
            return
        get = lambda r, f: text(ws.cell(r, cols[f]).value) if f in cols else None
        n = 0
        for r in range(hrow + 1, ws.max_row + 1):
            transcription, explanation, translation = get(r, 'transcription'), get(r, 'explanation'), get(r, 'translation')
            if not (transcription or explanation):
                continue
            try:
                key = ('id', str(int(float(get(r, 'id'))))) if get(r, 'id') else ('museum_id', get(r, 'museum_id'))
            except ValueError:
                self.warn.append(f'Transcriptions, row {r}: invalid Database ID "{get(r, "id")}", row skipped')
                continue
            if key not in self.touched:
                continue                                   # artwork not imported by this run
            aid, lang = self.touched[key]
            values = {'transcription': transcription, 'explanation': explanation, 'translation': translation}
            if lang and explanation:
                values[f'explanation_{lang}'] = explanation
            if self.cur.execute('SELECT 1 FROM transcriptions WHERE artwork_id = ?', (aid,)).fetchone():
                sets = ', '.join(f'{c} = ?' for c in values)
                self.cur.execute(f'UPDATE transcriptions SET {sets} WHERE artwork_id = ?', list(values.values()) + [aid])
            else:
                values['artwork_id'] = aid
                self.cur.execute(f'INSERT INTO transcriptions ({", ".join(values)}) VALUES ({", ".join("?" * len(values))})',
                                 list(values.values()))
            n += 1
        if n:
            self.report.append(f'transcriptions imported: {n}')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('db')
    ap.add_argument('xlsx')
    ap.add_argument('--assets', help='folder of the prepared media (e.g. ../assets)')
    ap.add_argument('--origin', help='only import the rows of this museum')
    ap.add_argument('--lang', choices=LANGS, help='language of the texts when the "Language" column is empty')
    ap.add_argument('--update', action='store_true', help='update artworks already in the database')
    ap.add_argument('--dry-run', action='store_true', help='check without writing anything')
    args = ap.parse_args()

    if not Path(args.db).is_file():
        sys.exit(f'database not found: {args.db}')
    wb = load_workbook(args.xlsx, data_only=True)
    if args.update and 'List of questions' in wb.sheetnames:
        sys.exit('PP1-Collection_Database.xlsx must not be used to update existing artworks: the database '
                 'has changed since (Le Fresnoy videos, merged types, corrections). Use --update only with '
                 'the spreadsheet of the museum concerned.')
    con = sqlite3.connect(args.db)
    imp = Importer(con, args)
    if imp.assets is None:
        imp.warn.append('--assets not given: media files were not checked')
    imp.import_artworks(pick_sheet(wb))
    imp.import_transcriptions(wb)

    total = con.execute('SELECT COUNT(*) FROM artworks').fetchone()[0]
    imp.report.append(f'collection: {total} artworks')
    no_vec = con.execute('SELECT COUNT(*) FROM artworks WHERE description_vector IS NULL').fetchone()[0]
    pending = con.execute('''SELECT COUNT(*) FROM artworks WHERE description_fr IS NULL
                             OR description_nl IS NULL OR description_en IS NULL''').fetchone()[0]
    if pending:
        imp.report.append(f'artworks to translate: {pending}')
    if no_vec:
        imp.report.append(f'artworks without a vector: {no_vec} (compute_embeddings.py, REQUIRED before '
                          'deploying: without vectors the kiosk server does not start)')

    if args.dry_run:
        con.rollback()
        imp.report.insert(0, 'DRY RUN (--dry-run): nothing was written')
    else:
        con.commit()
    con.close()
    print('\n'.join('  • ' + r for r in imp.report))
    if imp.mapping:
        print('    Museum ID → TRACES id')
        for m, a in imp.mapping:
            print(f'    {m or "—"} → {a}')
    for w in dict.fromkeys(imp.warn):
        print('  ⚠ ' + w)
    if not args.dry_run and imp.mapping:
        ids = [a for _m, a in imp.mapping]
        print(f'  → next step: python translate_collection.py {args.db} review_translations.xlsx '
              f'--ids {min(ids)}-{max(ids)} --dry-run')


if __name__ == '__main__':
    main()
